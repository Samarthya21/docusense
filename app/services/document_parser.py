from langchain_core.documents import Document
import pypdf
import docx
import os
import logging

logger = logging.getLogger(__name__)

def parse_pdf(file_path: str, filename: str) -> list[Document]:
    documents = []
    try:
        reader = pypdf.PdfReader(file_path)
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                # Page numbers are 1-indexed
                doc = Document(
                    page_content=text,
                    metadata={
                        "source": filename,
                        "page": page_idx + 1,
                        "file_type": "pdf"
                    }
                )
                documents.append(doc)
        logger.info(f"Successfully parsed PDF '{filename}' into {len(documents)} pages.")
    except Exception as e:
        logger.error(f"Error parsing PDF '{filename}': {e}")
        raise
    return documents

def parse_docx(file_path: str, filename: str) -> list[Document]:
    documents = []
    try:
        doc = docx.Document(file_path)
        current_section = "Introduction"
        current_text = []
        section_idx = 1
        
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            
            # Check if this paragraph looks like a heading
            is_heading = False
            if para.style.name.startswith("Heading"):
                is_heading = True
            elif len(text) < 100 and para.runs and all(r.bold for r in para.runs):
                is_heading = True
                
            if is_heading:
                # If we have accumulated paragraphs, flush them to a document chunk
                if current_text:
                    section_content = "\n".join(current_text)
                    documents.append(Document(
                        page_content=section_content,
                        metadata={
                            "source": filename,
                            "section": current_section,
                            "page": section_idx,  # Represent section index as 'page' for generic citation handling
                            "file_type": "docx"
                        }
                    ))
                    section_idx += 1
                    current_text = []
                current_section = text
            else:
                current_text.append(text)
                
        # Save any leftover text
        if current_text:
            section_content = "\n".join(current_text)
            documents.append(Document(
                page_content=section_content,
                metadata={
                    "source": filename,
                    "section": current_section,
                    "page": section_idx,
                    "file_type": "docx"
                }
            ))
            
        logger.info(f"Successfully parsed DOCX '{filename}' into {len(documents)} sections.")
    except Exception as e:
        logger.error(f"Error parsing DOCX '{filename}': {e}")
        raise
    return documents

def parse_document(file_path: str, filename: str) -> list[Document]:
    _, ext = os.path.splitext(filename.lower())
    if ext == ".pdf":
        return parse_pdf(file_path, filename)
    elif ext in [".docx", ".doc"]:
        return parse_docx(file_path, filename)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
